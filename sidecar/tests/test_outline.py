"""Syllabus outline extraction (slice 4): file→text + LLM extraction.

A fake provider stands in for Ollama so these are fast and deterministic.
"""

from __future__ import annotations

import pytest

from arbora_ai.llm.provider import LLMProvider, LLMSchemaError, LLMUnavailableError
from arbora_ai.outline.extract import extract_outline, extract_unit_info, syllabus_to_text
from arbora_ai.schemas.output import OutlineExtraction, UnitInfoExtraction

SYLLABUS = (
    "BIOL101 Cell Biology\n"
    "Coordinator: Dr Ada Chen (ada.chen@uni.edu)\n"
    "Lectures are online; tutorial attendance is a hurdle requirement.\n"
    "Week 1: Introduction to cells\n"
    "Week 2: The cell membrane\n"
    "Assignment 1 due 2026-03-15 (30%)\n"
    "Final exam 2026-06-20 (50%)\n"
)

UNIT_INFO_PAYLOAD = {
    "unit_code": "BIOL101",
    "coordinator_name": "Dr Ada Chen",
    "coordinator_contact": "ada.chen@uni.edu",
    "delivery_summary": "Online lectures with on-campus tutorials.",
    "classes": [
        {
            "label": "Tutorial",
            "schedule": "Wed 10:00-11:00",
            "mode": "on-campus",
            "attendance": "attendance is a hurdle requirement",
        }
    ],
    "assessments": [
        {"name": "Assignment 1", "weight_percent": 30, "due_text": "2026-03-15"},
        {"name": "Final exam", "weight_percent": 50, "due_text": "2026-06-20"},
    ],
}


class FakeProvider(LLMProvider):
    """Returns a schema-valid payload, branching on the requested schema."""

    def __init__(self, payload: dict | None = None):
        self._payload = payload

    def health(self) -> bool:
        return True

    def generate(self, prompt, schema=None, temperature=0.1):
        if self._payload is not None:
            return self._payload
        if schema and "unit_code" in schema.get("properties", {}):
            return UNIT_INFO_PAYLOAD
        return {
            "weeks": [
                {"week_number": 1, "title": "Introduction to cells", "summary": ""},
                {"week_number": 2, "title": "The cell membrane", "summary": ""},
            ],
            "deadlines": [
                {"title": "Assignment 1", "due_date": "2026-03-15", "type": "assignment"},
                {"title": "Final exam", "due_date": "2026-06-20", "type": "exam"},
            ],
        }


def test_extract_outline_parses_weeks_and_deadlines():
    result = extract_outline(FakeProvider(), SYLLABUS)
    assert isinstance(result, OutlineExtraction)
    assert [w.week_number for w in result.weeks] == [1, 2]
    assert result.weeks[1].title == "The cell membrane"
    assert result.deadlines[0].type == "assignment"
    assert result.deadlines[1].due_date == "2026-06-20"


def test_extract_outline_retries_then_raises_on_junk():
    # Never schema-valid (week_number 0 < 1) → exhausts retries and raises.
    bad = FakeProvider({"weeks": [{"week_number": 0, "title": "x"}], "deadlines": []})
    with pytest.raises(LLMSchemaError):
        extract_outline(bad, SYLLABUS)


def test_extract_outline_accepts_empty_arrays():
    empty = FakeProvider({"weeks": [], "deadlines": []})
    result = extract_outline(empty, SYLLABUS)
    assert result.weeks == []
    assert result.deadlines == []


def test_extract_unit_info_parses_all_sections():
    info = extract_unit_info(FakeProvider(), SYLLABUS)
    assert info.unit_code == "BIOL101"
    assert info.coordinator_contact == "ada.chen@uni.edu"
    assert info.classes[0].attendance == "attendance is a hurdle requirement"
    assert info.assessments[0].weight_percent == 30
    assert info.assessments[1].due_text == "2026-06-20"


def test_extract_unit_info_degrades_to_empty_on_junk():
    # Never schema-valid (weight 200 > 100) → empty result, no raise.
    bad = FakeProvider(
        {"assessments": [{"name": "A1", "weight_percent": 200}], "classes": []}
    )
    info = extract_unit_info(bad, SYLLABUS)
    assert info == UnitInfoExtraction()


def test_extract_unit_info_degrades_to_empty_when_provider_dies():
    class DeadProvider(LLMProvider):
        def health(self) -> bool:
            return False

        def generate(self, prompt, schema=None, temperature=0.1):
            raise LLMUnavailableError("gone")

    info = extract_unit_info(DeadProvider(), SYLLABUS)
    assert info == UnitInfoExtraction()


def test_syllabus_to_text_reads_text_file(tmp_path):
    f = tmp_path / "syllabus.txt"
    f.write_text(SYLLABUS, encoding="utf-8")
    assert "cell membrane" in syllabus_to_text(f)


def test_syllabus_to_text_rejects_unsupported(tmp_path):
    bad = tmp_path / "syllabus.csv"
    bad.write_text("week,topic\n1,cells", encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported"):
        syllabus_to_text(bad)


def test_syllabus_to_text_empty_file(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError, match="no text could be extracted"):
        syllabus_to_text(empty)


def test_syllabus_to_text_reads_docx(tmp_path):
    from docx import Document

    path = tmp_path / "syllabus.docx"
    doc = Document()
    doc.add_paragraph("Week 1: Introduction to cells")
    doc.add_paragraph("Week 2: The cell membrane")
    doc.save(str(path))

    text = syllabus_to_text(path)
    assert "Introduction to cells" in text
    assert "cell membrane" in text


def test_syllabus_to_text_keeps_back_page_schedule(tmp_path):
    """Regression: institutional unit guides bury the weekly-activities table
    and assessment schedule on the LAST page, behind a title page + table of
    contents. An over-budget syllabus must keep those, not the boilerplate head."""
    toc = "\n\n".join(f"Section {i} .................... {i}" for i in range(400))  # padding
    schedule = (
        "Unit weekly activities\n"
        "Week 1: What is deep learning?\n"
        "Week 6: Computer vision\n"
        "Assessment task 1 due 24 July 2026 (30%)\n"
        "Final exam due 5 October 2026 (40%)\n"
    )
    path = tmp_path / "unit-guide.txt"
    path.write_text(f"SIT319 Deep Learning\n\n{toc}\n\n{schedule}", encoding="utf-8")

    text = syllabus_to_text(path)
    assert len(text) <= 12000  # over-budget input was reduced
    assert "Unit weekly activities" in text  # ...but the back-page schedule survived
    assert "Assessment task 1 due 24 July 2026" in text
    assert "Final exam due 5 October 2026" in text
