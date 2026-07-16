from pathlib import Path

import pytest

from arbora_ai.ingest.chunk import chunk_units
from arbora_ai.ingest.figures import extract_pdf_figures
from arbora_ai.ingest.parse import (
    SourceUnit,
    detect_type,
    parse_docx,
    parse_pdf,
    parse_pptx,
    parse_text,
)


def _solid_png(w: int, h: int) -> bytes:
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.clear_with(180)
    return pix.tobytes("png")


def _pdf_with_images(path: Path, sizes: list[tuple[int, int]]) -> None:
    """A one-page PDF with an embedded image for each (w, h) in `sizes`."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=800, height=1000)
    y = 20
    for w, h in sizes:
        page.insert_image(pymupdf.Rect(20, y, 20 + w, y + h), stream=_solid_png(w, h))
        y += h + 20
    doc.save(str(path))
    doc.close()


def test_extract_pdf_figures_keeps_real_drops_chrome(tmp_path):
    pdf = tmp_path / "src.pdf"
    _pdf_with_images(pdf, [(400, 300), (40, 40)])  # one real figure, one icon

    figs = extract_pdf_figures(pdf, tmp_path / "out", "src-1")

    assert len(figs) == 1  # the 40x40 icon is filtered out
    fig = figs[0]
    assert fig.page == 1
    assert (fig.width, fig.height) == (400, 300)
    assert Path(fig.path).exists()
    assert Path(fig.path).parent.name == "src-1"


def test_extract_pdf_figures_none_when_no_qualifying_images(tmp_path):
    pdf = tmp_path / "empty.pdf"
    _pdf_with_images(pdf, [(30, 30)])  # only a tiny image
    assert extract_pdf_figures(pdf, tmp_path / "out", "s") == []


@pytest.mark.parametrize("ext", ["mp3", "m4a", "wav", "ogg", "flac", "aac", "mp4", "mkv", "webm", "avi", "mov"])
def test_detect_type_audio_and_video(ext):
    assert detect_type(f"lecture.{ext}") == "audio"


def test_detect_type_pdf_and_slide():
    assert detect_type("notes.pdf") == "pdf"
    assert detect_type("deck.pptx") == "slide"


def test_detect_type_doc_and_text():
    assert detect_type("essay.docx") == "doc"
    assert detect_type("notes.txt") == "text"
    assert detect_type("README.md") == "text"
    assert detect_type("guide.markdown") == "text"


def test_detect_type_rejects_unknown():
    with pytest.raises(ValueError, match="unsupported"):
        detect_type("archive.zip")


def test_parse_text_splits_paragraphs(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text(
        "# Cell biology\n\nThe mitochondrion is the powerhouse of the cell.\n\n\nRibosomes synthesise proteins.\n",
        encoding="utf-8",
    )

    units = parse_text(path)
    assert len(units) == 3
    assert units[0].location == {"type": "page", "page": 1}
    assert "mitochondrion" in units[1].text
    assert units[2].location == {"type": "page", "page": 3}


def test_parse_docx_reads_paragraphs(tmp_path):
    from docx import Document

    path = tmp_path / "syllabus.docx"
    doc = Document()
    doc.add_paragraph("Week 1: Introduction to Biology")
    doc.add_paragraph("Week 2: Cell Structure and Function")
    doc.add_paragraph("")  # blank paragraph — should be skipped
    doc.add_paragraph("Final exam: June 20")
    doc.save(str(path))

    units = parse_docx(path)
    assert len(units) == 3
    assert units[0].location == {"type": "page", "page": 1}
    assert "Introduction" in units[0].text
    assert units[2].location == {"type": "page", "page": 3}


def test_parse_docx_reads_table_rows(tmp_path):
    """Regression: syllabuses put the weekly schedule in a table, and
    `doc.paragraphs` skips table cells — so table rows must survive parsing."""
    from docx import Document

    path = tmp_path / "syllabus.docx"
    doc = Document()
    doc.add_paragraph("BIOL101 Cell Biology")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "1"
    table.cell(0, 1).text = "Introduction to cells"
    table.cell(1, 0).text = "2"
    table.cell(1, 1).text = "The cell membrane"
    doc.save(str(path))

    text = "\n".join(u.text for u in parse_docx(path))
    assert "Introduction to cells" in text
    assert "The cell membrane" in text


def test_parse_pptx_reads_table_rows(tmp_path):
    """Regression: a schedule slide is often a bare table with no text frame —
    its rows must be extracted, not silently dropped."""
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "1"
    table.cell(0, 1).text = "Introduction to cells"
    table.cell(1, 0).text = "2"
    table.cell(1, 1).text = "The cell membrane"
    prs.save(str(path))

    text = "\n".join(u.text for u in parse_pptx(path))
    assert "Introduction to cells" in text
    assert "The cell membrane" in text


def test_chunk_keeps_location_and_splits_long_text():
    long = " ".join(f"Sentence number {i} about cells." for i in range(60))
    units = [SourceUnit(text=long, location={"type": "page", "page": 2})]
    chunks = chunk_units(units, source_id="bio", max_chars=200)
    assert len(chunks) > 1
    assert all(c.location == {"type": "page", "page": 2} for c in chunks)
    assert all(c.source_id == "bio" for c in chunks)
    assert all(len(c.text) <= 200 for c in chunks)


def test_chunk_drops_tiny_fragments():
    units = [SourceUnit(text="Title", location={"type": "page", "page": 1})]
    assert chunk_units(units, source_id="s", min_chars=40) == []


def test_chunk_prefers_paragraph_boundaries():
    para1 = "First paragraph sentence one. First paragraph sentence two."
    para2 = "Second paragraph about another idea entirely."
    units = [SourceUnit(text=f"{para1}\n\n{para2}", location={"type": "page", "page": 1})]
    chunks = chunk_units(units, source_id="s", max_chars=80, min_chars=10)
    texts = [c.text for c in chunks]
    # split lands on the blank line, not mid-paragraph
    assert texts == [para1, para2]


def test_chunk_overlap_repeats_boundary_sentence():
    sentences = [f"Sentence number {i} about photosynthesis." for i in range(10)]
    units = [SourceUnit(text=" ".join(sentences), location={"type": "page", "page": 1})]
    chunks = chunk_units(units, source_id="s", max_chars=120, min_chars=10)
    assert len(chunks) > 1
    assert all(len(c.text) <= 120 for c in chunks)
    # consecutive pieces share their boundary sentence, so an idea cut at the
    # boundary is retrievable from either side
    for prev, cur in zip(chunks, chunks[1:]):
        last_sentence = prev.text.rsplit(". ", 1)[-1]
        assert cur.text.startswith(last_sentence.split(".")[0])


def test_parse_pdf_reads_pages_with_numbers(tmp_path):
    import pymupdf

    path = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"This is page {i + 1} about photosynthesis and chloroplasts.")
    doc.save(str(path))
    doc.close()

    units = parse_pdf(path)
    assert len(units) == 3
    assert units[0].location == {"type": "page", "page": 1}
    assert "page 1" in units[0].text
    assert units[2].location["page"] == 3
