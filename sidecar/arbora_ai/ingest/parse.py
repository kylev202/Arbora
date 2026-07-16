"""Parse a source file into located text units.

A `SourceUnit` is one citable region of a document — a PDF page, a slide, or an
audio segment — paired with the location that a citation will point back to.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class SourceUnit:
    """One citable region of a source. `location` matches the SourceRef schema:
    {"type": "page", "page": N} or {"type": "timestamp", "timestamp_ms": N}."""

    text: str
    location: dict


def parse_pdf(path: str | Path) -> list[SourceUnit]:
    """One unit per page (1-based page numbers)."""
    import pymupdf

    units: list[SourceUnit] = []
    with pymupdf.open(str(path)) as doc:
        for i, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                units.append(SourceUnit(text=text, location={"type": "page", "page": i + 1}))
    return units


def parse_pptx(path: str | Path) -> list[SourceUnit]:
    """One unit per slide (1-based; slides cite as `page` per the IPC contract)."""
    from pptx import Presentation

    units: list[SourceUnit] = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif shape.has_table:
                # Schedule slides are often a bare table with no text frame —
                # without this the whole slide would come back empty.
                for row in shape.table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells)
                    if line.strip(" |"):
                        parts.append(line)
        text = "\n".join(parts).strip()
        if text:
            units.append(SourceUnit(text=text, location={"type": "page", "page": i + 1}))
    return units


def _iter_docx_blocks(doc):
    """Yield paragraph and table text in document order.

    `doc.paragraphs` skips everything inside tables, so a syllabus that lays its
    weekly schedule or assessment breakdown out in a table (most do) would lose
    exactly that. Walk the body children instead: paragraphs verbatim, table
    rows as pipe-joined cells so row/column structure survives as plain text."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc).text
        elif isinstance(child, CT_Tbl):
            for row in Table(child, doc).rows:
                yield " | ".join(cell.text.strip() for cell in row.cells)


def parse_docx(path: str | Path) -> list[SourceUnit]:
    """One unit per non-empty paragraph or table row, numbered sequentially
    (cited as pages). Table rows are included — see `_iter_docx_blocks`."""
    from docx import Document

    doc = Document(str(path))
    units: list[SourceUnit] = []
    page = 1
    for block in _iter_docx_blocks(doc):
        text = block.strip()
        if text.strip(" |"):
            units.append(SourceUnit(text=text, location={"type": "page", "page": page}))
            page += 1
    return units


def parse_text(path: str | Path) -> list[SourceUnit]:
    """One unit per non-empty paragraph (blank-line separated), numbered
    sequentially. Plain text has no real pages, so paragraphs are cited as
    pages — keeping every chunk traceable (law #1)."""
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    units: list[SourceUnit] = []
    page = 1
    for block in re.split(r"\n\s*\n", raw):
        text = block.strip()
        if text:
            units.append(SourceUnit(text=text, location={"type": "page", "page": page}))
            page += 1
    return units


def parse_audio(path: str | Path, transcriber=None) -> list[SourceUnit]:
    """One unit per transcript segment, located by start timestamp (ms)."""
    if transcriber is None:
        from ..transcribe.whisper import Transcriber

        transcriber = Transcriber()
    segments = transcriber.transcribe(str(path))
    return [
        SourceUnit(text=seg.text, location={"type": "timestamp", "timestamp_ms": seg.start_ms})
        for seg in segments
        if seg.text.strip()
    ]


_PDF = {".pdf"}
_SLIDE = {".pptx", ".ppt"}
_AUDIO = {".mp3", ".m4a", ".wav", ".ogg", ".flac", ".aac", ".mp4", ".mkv", ".webm", ".avi", ".mov"}
_DOC = {".docx"}
_TEXT = {".txt", ".md", ".markdown"}


def detect_type(path: str | Path) -> str:
    """Map a file extension to a source type
    ("pdf" | "slide" | "audio" | "doc" | "text")."""
    ext = Path(path).suffix.lower()
    if ext in _PDF:
        return "pdf"
    if ext in _SLIDE:
        return "slide"
    if ext in _AUDIO:
        return "audio"
    if ext in _DOC:
        return "doc"
    if ext in _TEXT:
        return "text"
    raise ValueError(f"unsupported file type: {ext!r}")


def parse_source(path: str | Path, source_type: str | None = None, transcriber=None) -> list[SourceUnit]:
    """Parse any supported source into located units (type auto-detected)."""
    source_type = source_type or detect_type(path)
    if source_type == "pdf":
        return parse_pdf(path)
    if source_type == "slide":
        return parse_pptx(path)
    if source_type == "audio":
        return parse_audio(path, transcriber=transcriber)
    if source_type == "doc":
        return parse_docx(path)
    if source_type == "text":
        return parse_text(path)
    raise ValueError(f"unknown source type: {source_type!r}")
